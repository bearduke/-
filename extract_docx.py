# -*- coding: utf-8 -*-
"""
读取 Word 文档并提取：
1. 完整文本内容（逐段）
2. 文档结构（标题、段落、表格、字体大小、加粗/斜体）
3. 表格及其内容
"""
import os
from docx import Document
from docx.shared import Pt

DOC_PATH = r"c:\Users\monglongshi\Desktop\执行小助手精简版v1\115_利息计算说明.docx"


def style_name_to_type(style_name):
    """根据样式名判断类型"""
    if not style_name:
        return "正文"
    s = str(style_name).lower()
    if "heading 1" in s or s in ("heading 1", "标题 1", "标题1"):
        return "标题1"
    if "heading 2" in s or s in ("heading 2", "标题 2", "标题2"):
        return "标题2"
    if "heading 3" in s or s in ("heading 3", "标题 3", "标题3"):
        return "标题3"
    if "heading" in s or "标题" in s:
        return f"标题({style_name})"
    if "title" in s or s == "标题":
        return "文档标题"
    if "list" in s or "列表" in s:
        return f"列表({style_name})"
    return f"正文({style_name})"


def get_font_info(run):
    """获取 run 的字体信息"""
    info = {}
    # 字体名
    font_name = run.font.name
    # 东亚字体（中文字体）
    east_asia = None
    try:
        rpr = run._element.rPr
        if rpr is not None:
            rfonts = rpr.rFonts
            if rfonts is not None:
                east_asia = rfonts.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"
                )
    except Exception:
        pass
    # 字号
    size_pt = None
    if run.font.size is not None:
        size_pt = run.font.size.pt
    info["font_name"] = font_name
    info["east_asia_font"] = east_asia
    info["size_pt"] = size_pt
    info["bold"] = run.font.bold
    info["italic"] = run.font.italic
    info["underline"] = run.font.underline
    info["color"] = None
    try:
        if run.font.color and run.font.color.rgb is not None:
            info["color"] = str(run.font.color.rgb)
    except Exception:
        pass
    return info


def main():
    if not os.path.exists(DOC_PATH):
        print(f"错误：文件不存在 -> {DOC_PATH}")
        return

    doc = Document(DOC_PATH)

    print("=" * 80)
    print("文档路径:", DOC_PATH)
    print("=" * 80)

    # 核心属性
    print("\n【文档核心属性】")
    cp = doc.core_properties
    for attr in ["author", "title", "subject", "created", "modified",
                 "last_modified_by", "category", "comments", "keywords"]:
        val = getattr(cp, attr, None)
        if val:
            print(f"  {attr}: {val}")

    # 段落级元素遍历（按文档顺序）
    print("\n" + "=" * 80)
    print("一、文档完整内容（按文档顺序，逐段输出）")
    print("=" * 80)

    body = doc.element.body
    # 建立 element -> paragraph / table 的映射
    para_map = {p._element: p for p in doc.paragraphs}
    table_map = {t._element: t for t in doc.tables}

    para_idx = 0
    table_idx = 0
    for child in body.iterchildren():
        if child in para_map:
            p = para_map[child]
            para_idx += 1
            style_name = p.style.name if p.style else None
            style_type = style_name_to_type(style_name)
            text = p.text
            print(f"\n[段落 #{para_idx}] 类型={style_type} | 原样式名={style_name}")
            print(f"  对齐方式: {p.alignment}")
            print(f"  文本: {text if text else '(空段落)'}")
            # 逐 run 输出字体信息
            if p.runs:
                for ri, run in enumerate(p.runs):
                    if not run.text.strip():
                        continue
                    fi = get_font_info(run)
                    print(f"    Run#{ri}: '{run.text}'")
                    print(f"      字体(西文)={fi['font_name']}, 字体(中文)={fi['east_asia_font']}, "
                          f"字号={fi['size_pt']}pt, 加粗={fi['bold']}, 斜体={fi['italic']}, "
                          f"下划线={fi['underline']}, 颜色={fi['color']}")
        elif child in table_map:
            t = table_map[child]
            table_idx += 1
            print("\n" + "-" * 60)
            print(f"[表格 #{table_idx}]")
            print(f"  行数: {len(t.rows)}, 列数: {len(t.columns)}")
            print(f"  表格样式: {t.style.name if t.style else None}")
            print("-" * 60)
            for ri, row in enumerate(t.rows):
                cells_text = []
                for ci, cell in enumerate(row.cells):
                    cell_text = cell.text.replace("\n", " / ")
                    cells_text.append(cell_text)
                print(f"  行{ri}: {cells_text}")
            # 表格内单元格字体信息
            print("  --- 表格单元格字体信息 ---")
            for ri, row in enumerate(t.rows):
                for ci, cell in enumerate(row.cells):
                    for pi, p in enumerate(cell.paragraphs):
                        for rri, run in enumerate(p.runs):
                            if not run.text.strip():
                                continue
                            fi = get_font_info(run)
                            print(f"    [行{ri},列{ci}] '{run.text}' -> "
                                  f"字号={fi['size_pt']}pt, 加粗={fi['bold']}, "
                                  f"斜体={fi['italic']}, 中文字体={fi['east_asia_font']}, "
                                  f"西文字体={fi['font_name']}")
            print("-" * 60)

    # 汇总
    print("\n" + "=" * 80)
    print("二、文档结构汇总")
    print("=" * 80)
    print(f"  段落总数: {len(doc.paragraphs)}")
    print(f"  表格总数: {len(doc.tables)}")

    # 标题列表
    print("\n  标题列表:")
    heading_count = 0
    for i, p in enumerate(doc.paragraphs, 1):
        sn = p.style.name if p.style else ""
        if sn and ("heading" in sn.lower() or "标题" in sn or sn.lower() == "title"):
            heading_count += 1
            print(f"    段落#{i} [{style_name_to_type(sn)}]: {p.text}")
    if heading_count == 0:
        print("    (未检测到标准标题样式)")

    # 节信息
    print(f"\n  节(Section)数: {len(doc.sections)}")
    for si, sec in enumerate(doc.sections, 1):
        print(f"    节{si}: 页面宽度={sec.page_width}, 页面高度={sec.page_height}")
        print(f"          左边距={sec.left_margin}, 右边距={sec.right_margin}, "
              f"上边距={sec.top_margin}, 下边距={sec.bottom_margin}")

    print("\n提取完成。")


if __name__ == "__main__":
    main()
